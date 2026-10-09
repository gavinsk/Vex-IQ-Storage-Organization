"""Build tools/bin-assignments.xlsx from data/parts.json.

One row per bin; type part names into the Part columns (Excel suggests matches
as you type). Apply the sheet back with tools/apply_bin_sheet.py.

Needs openpyxl (pip install openpyxl). Regenerating overwrites the sheet with
whatever parts.json currently says, so apply any edits first.
"""
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tools" / "bin-assignments.xlsx"
PART_COLS = 5          # Part 1 … Part 5
SPARE_ZONE_IDS = 4     # empty V/E rows beyond the highest one in use

parts = json.loads((ROOT / "data/parts.json").read_text())
layout = json.loads((ROOT / "data/storage-layout.json").read_text())

by_bin = defaultdict(list)
for p in parts:
    for loc in [p["location"], *p.get("extraLocations", [])]:
        by_bin[loc].append(p)

# ---- every bin, in physical order -------------------------------------------
bins = []  # (id, zone)
for side in layout["cart"]["sides"]:
    for g in layout["cart"]["binGroups"]:
        size = "Small" if g["firstRow"] == 1 else "Medium"
        for r in range(g["firstRow"], g["lastRow"] + 1):
            for c in range(1, g["cols"] + 1):
                bins.append((f"{side['id']}{r}-{c}", f"Cart {side['id']} – {size}"))
ak = layout["akroMils"]
for n in range(1, ak["drawers"] + 1):
    bins.append((f"{ak['prefix']}{n:02d}", "Akro-Mils"))
for z in layout["zones"]:
    used = [int(b[len(z["prefix"]):]) for b in by_bin if b.startswith(z["prefix"]) and b[len(z["prefix"]):].isdigit()]
    for n in range(1, max(used, default=0) + SPARE_ZONE_IDS + 1):
        bins.append((f"{z['prefix']}{n:02d}", z["name"]))

known = {b for b, _ in bins}
missing = [p["name"] for p in parts if not {p["location"], *p.get("extraLocations", [])} <= known]
if missing:
    raise SystemExit(f"Parts with locations not in the sheet: {missing}")

# ---- styles -----------------------------------------------------------------
head = Font(bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="1F4E79")
zone_fills = {
    "Cart A": PatternFill("solid", fgColor="E8F0FA"),
    "Cart B": PatternFill("solid", fgColor="EAF6EA"),
    "Akro-Mils": PatternFill("solid", fgColor="FFF4E0"),
}
other_fill = PatternFill("solid", fgColor="F2F2F2")

wb = Workbook()

# ---- Bins sheet -------------------------------------------------------------
ws = wb.active
ws.title = "Bins"
headers = ["Bin", "Zone", "Bin Label"] + [f"Part {i}" for i in range(1, PART_COLS + 1)]
ws.append(headers)
for cell in ws[1]:
    cell.font, cell.fill = head, head_fill
for bin_id, zone in bins:
    contents = by_bin.get(bin_id, [])
    label = contents[0]["binLabel"] if contents else ""
    ws.append([bin_id, zone, label] + [p["name"] for p in contents])
    fill = next((f for k, f in zone_fills.items() if zone.startswith(k)), other_fill)
    for col in (1, 2):
        ws.cell(ws.max_row, col).fill = fill
    ws.cell(ws.max_row, 1).font = Font(bold=True)
last_row = ws.max_row
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{chr(64 + len(headers))}{last_row}"
ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 34
for i in range(PART_COLS):
    ws.column_dimensions[chr(68 + i)].width = 30

part_range = f"D2:{chr(67 + PART_COLS)}{last_row}"
dv = DataValidation(
    type="list",
    formula1=f"=Parts!$A$2:$A${len(parts) + 1}",
    allow_blank=True,
    showErrorMessage=True,
    errorTitle="Not a known part",
    error="Pick a part name from the suggestions (see the Parts tab).",
)
dv.add(part_range)
ws.add_data_validation(dv)

# A part typed into more than one bin turns light blue (fine for high-volume parts; a typo otherwise).
ws.conditional_formatting.add(
    part_range,
    FormulaRule(
        formula=[f"AND(D2<>\"\",COUNTIF($D$2:${chr(67 + PART_COLS)}${last_row},D2)>1)"],
        fill=PatternFill("solid", fgColor="CFE2F3"),
    ),
)

# ---- Parts sheet (suggestion list + where each part is) ---------------------
ps = wb.create_sheet("Parts")
ps.append(["Part name", "Category", "In bin(s)", "Bins", "Part numbers"])
for cell in ps[1]:
    cell.font, cell.fill = head, head_fill
bins_range = f"Bins!$D$2:${chr(67 + PART_COLS)}${last_row}"
for i, p in enumerate(parts, start=2):
    ps.append([
        p["name"],
        p["category"],
        None,  # "In bin(s)": array formula set below
        f"=COUNTIF({bins_range},A{i})",
        ", ".join(p["partNumbers"]),
    ])
# Every bin row that lists this part, e.g. "A3-5, A3-6".
any_col = "+".join(f"(Bins!${c}$2:${c}${last_row}=A{{i}})" for c in (chr(68 + k) for k in range(PART_COLS)))
for i in range(2, len(parts) + 2):
    ps[f"C{i}"] = ArrayFormula(f"C{i}", f'=TEXTJOIN(", ",TRUE,IF({any_col.format(i=i)},Bins!$A$2:$A${last_row},""))')
ps.freeze_panes = "B2"
ps.auto_filter.ref = f"A1:E{len(parts) + 1}"
for col, w in zip("ABCDE", (34, 18, 16, 7, 40)):
    ps.column_dimensions[col].width = w
for i in range(2, len(parts) + 2):
    ps.cell(i, 3).alignment = Alignment(horizontal="center")
    ps.cell(i, 4).alignment = Alignment(horizontal="center")
count_range = f"D2:D{len(parts) + 1}"
ps.conditional_formatting.add(count_range, FormulaRule(formula=["D2=0"], fill=PatternFill("solid", fgColor="FF9999")))
ps.conditional_formatting.add(count_range, FormulaRule(formula=["D2>1"], fill=PatternFill("solid", fgColor="CFE2F3")))

# ---- How-to sheet -----------------------------------------------------------
hs = wb.create_sheet("How to use")
for line in [
    "Bin assignments — editing sheet for data/parts.json",
    "",
    "Bins tab: one row per bin. Type in a Part column; Excel suggests matching part names (arrow keys + Enter to pick).",
    "Bin Label: what the bin's label/website should say. Leave blank to use the part names.",
    "To move a part: delete it from its old bin and type it into the new one.",
    "A part can go in more than one bin (high-volume parts): type it in each bin. The first bin is its main one.",
    "Light blue cell = that part is in more than one bin — check it's on purpose. Parts tab: red = not placed anywhere.",
    "Parts that share a bin should share one Bin Label; a part split across bins uses the first bin's label.",
    "",
    "When done: save, then run   python3 tools/apply_bin_sheet.py   (preview)",
    "and         python3 tools/apply_bin_sheet.py --apply   (writes data/parts.json).",
    "",
    "Sheet is generated from parts.json by tools/make_bin_sheet.py — regenerating overwrites unapplied edits.",
]:
    hs.append([line])
hs["A1"].font = Font(bold=True, size=14)
hs.column_dimensions["A"].width = 110

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT.relative_to(ROOT)}: {len(bins)} bins, {len(parts)} parts")
